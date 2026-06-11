from flask import Flask, request, jsonify, send_from_directory, session
from flask_cors import CORS
import os
import sqlite3
import win32com.client
import pythoncom
import time
import logging
import openpyxl
from functools import wraps

# --- Configure Logging ---
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

# Resolve path for a single-folder setup
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FRONTEND_DIR = BASE_DIR 

app = Flask(__name__, static_folder=FRONTEND_DIR, static_url_path='')
app.secret_key = "kirloskar_super_secret_key_2025" # Required to use sessions/login
CORS(app)

# Define Valid Users (Username: Password)
USERS = {
    "admin": "password123",
    "user": "kirloskar2025"
}

# ==================== AUTHENTICATION HELPERS ====================

def verify_user(username, password):
    """Checks if the username and password match our records."""
    return USERS.get(username) == password

def login_required(f):
    """
    Decorator to protect routes. 
    Checks if 'username' is in the session.
    """
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'username' not in session:
            return jsonify({"error": "Unauthorized. Please login first."}), 401
        return f(*args, **kwargs)
    return decorated_function

# ==================== EXCEL FILES CONFIGURATION ====================
EXCEL_FILES = {
    "kwk": "water.xlsm",
    "kwi": "kwi.xlsm",
    "kws": "kws.xlsm",
    "kaf": "output.xlsm",
    "kaa": "kaa.xlsm",
    "kas": "kas.xlsm"
}

SHEETS = ["Input data", "Sheet3", "Sheet4"]

TARGET_CELLS = {
    "Input data": {
        "water": ["F13", "F12", "F17", "F18", "F16", "J12", "J13", "J17", "J18", "J16", "F20", "J20"],
        "air":   ["F13", "F12", "F17", "F18", "F16", "J12", "J13", "F20"]
    },
    "Sheet3": ["AF14"]
}

SHEET3_ADJUSTED_CELLS = {
    "s_value_cell": "AG9",
    "d_value_cell": "AG10"
}

AIR_COOLED_CELLS = {
    "ambient_temp": "AG11",
    "altitude":     "AG6",
    "fin_material": "AG7",
    "wet_bulb_temp":"AG8"
}

RESULT_COLUMN_CELLS = {
    "Sheet4": {
        "B_column": [f"B{i}" for i in range(9, 35)],
        "C_column": [f"C{i}" for i in range(9, 35)],
        "D_column": [f"D{i}" for i in range(9, 35)],
        "E_column": [f"E{i}" for i in range(9, 37)],
    },
    "Input data": {
        "N_column": [f"N{i}" for i in range(11, 17)]
    }
}

EVAP_LEAVING_TEMP_INDEX    = 1
COND_LEAVING_TEMP_INDEX    = 6
REQUIRED_CAPACITY_INDEX    = 12
CHILLER_TYPE_INDEX         = 13
AMBIENT_TEMP_INDEX         = 14
ALTITUDE_INDEX             = 15
FIN_MATERIAL_INDEX         = 16
WET_BULB_TEMP_INDEX        = 17
KW_TO_TONR_FACTOR          = 0.284345

def get_ws_by_name(workbook, name):
    try:
        return workbook.Worksheets(name)
    except Exception:
        return workbook.Sheets(name)

def safe_val(v, default=""):
    return default if v is None else v

# ==================== GA & PID EXCEL PATH ====================
GAPID_EXCEL_PATH = os.path.join(BASE_DIR, 'SO, GA, and PID Sheet.xlsx')


# ==================== SERVE FRONTEND ====================
@app.route('/')
def serve_root():
    return send_from_directory(FRONTEND_DIR, 'login.html')

@app.route('/<path:filename>')
def serve_frontend(filename):
    return send_from_directory(FRONTEND_DIR, filename)


# ==================== API ROUTES ====================
@app.route('/login', methods=['POST'])
def login():
    data     = request.get_json()
    username = data.get('username')
    password = data.get('password')
    
    if verify_user(username, password):
        # Store user in session
        session['username'] = username
        logging.info(f"User '{username}' logged in.")
        return jsonify({"message": "Login successful", "username": username}), 200
    
    return jsonify({"error": "Invalid username or password"}), 401

@app.route('/logout', methods=['POST'])
def logout():
    session.pop('username', None)
    return jsonify({"message": "Logged out successfully"}), 200

@app.route('/status', methods=['GET'])
def status():
    return jsonify({"status": "online"}), 200


# ==================== DOCUMENT & PDF ROUTES ====================

@app.route('/search-gapid', methods=['GET'])
@login_required
def search_gapid():
    query       = request.args.get('query', '').strip()
    search_type = request.args.get('search_type', 'so').strip().lower()

    if not query:
        return jsonify({"found": False, "error": "No query provided"}), 400

    if not os.path.exists(GAPID_EXCEL_PATH):
        return jsonify({"found": False, "error": f"Excel file not found on server."}), 404

    try:
        wb = openpyxl.load_workbook(GAPID_EXCEL_PATH, read_only=True, data_only=True)
        ws = wb['DATA']   

        for row_num in range(4, (ws.max_row or 4) + 1):
            so_raw    = ws[f'C{row_num}'].value
            model_raw = ws[f'E{row_num}'].value
            link_raw  = ws[f'G{row_num}'].value

            if so_raw is None: continue

            so_val = str(int(so_raw)) if isinstance(so_raw, float) else str(so_raw).strip()
            model_val = str(model_raw).strip() if model_raw is not None else ''
            link_val  = str(link_raw).strip()  if link_raw  is not None else ''

            matched = False
            if search_type == 'so':
                parts = [p.strip() for p in so_val.replace('/', ',').split(',')]
                if query == so_val or query in parts or (parts and query == parts[0]):
                    matched = True
            else: # search by model
                matched = (query.lower() == model_val.lower())

            if matched:
                customer_val = str(ws[f'D{row_num}'].value or '')
                wb.close()
                
                if not link_val.startswith('http'):
                    return jsonify({"found": False, "error": "Record found but no valid link attached."}), 404
                
                return jsonify({
                    "found":    True,
                    "so":       so_val,
                    "model":    model_val,
                    "customer": customer_val,
                    "link":     link_val
                }), 200

        wb.close()
        return jsonify({"found": False, "error": f"No record found for '{query}'."}), 404

    except Exception as e:
        logging.error(f"Error searching GA/PID sheet: {e}")
        return jsonify({"found": False, "error": f"Server error: {str(e)}"}), 500


ALLOWED_FOLDERS = [
    'ENGINEERING BULLETIN', 
    'Engg. Specification', 
    'tds data', 
    'Product Info', 
    'Certificates'
]

@app.route('/api/document/<folder>/<path:filename>', methods=['GET'])
def get_document(folder, filename):
    if folder not in ALLOWED_FOLDERS:
        return f"Invalid folder request: {folder}", 400

    target_dir = os.path.join(BASE_DIR, folder)
    file_path = os.path.join(target_dir, filename)
    
    if os.path.exists(file_path):
        return send_from_directory(target_dir, filename, as_attachment=False)
    else:
        return f"File '{filename}' not found.", 404


@app.route('/api/download-document/<folder>/<path:filename>', methods=['GET'])
def download_general_doc(folder, filename):
    if folder not in ALLOWED_FOLDERS:
        return f"Invalid folder request: {folder}", 400

    target_dir = os.path.join(BASE_DIR, folder)
    file_path = os.path.join(target_dir, filename)
    
    if os.path.exists(file_path):
        return send_from_directory(target_dir, filename, as_attachment=True)
    else:
        return f"File '{filename}' not found.", 404


# ==================== EXCEL CALCULATION ROUTE ====================
@app.route('/write-excel', methods=['POST'])
@login_required
def write_excel():
    excel = None
    workbook = None
    pythoncom.CoInitialize()

    response_data = { "message": "Processing started.", "results": {}, "detailed_results": {}, "error": None }

    try:
        request_data = request.get_json()
        values = list(request_data.get('values', []))
        chiller_type   = request_data.get('chillerType', '').lower()
        chiller_series = request_data.get('chillerSeries', '').lower()

        if not chiller_type and len(values) > CHILLER_TYPE_INDEX:
            chiller_type = str(values[CHILLER_TYPE_INDEX]).lower()
        if not chiller_series:
            chiller_series = "kwk" if chiller_type == "water" else "kaf"

        if chiller_type not in ["water", "air"]:
            response_data["error"] = f"Invalid base type: {chiller_type}"
            return jsonify(response_data), 400
        if chiller_series not in EXCEL_FILES:
            response_data["error"] = f"Invalid chiller series: {chiller_series}"
            return jsonify(response_data), 400

        excel_filename = EXCEL_FILES[chiller_series]
        abs_path = os.path.join(BASE_DIR, excel_filename)

        if not os.path.exists(abs_path):
            response_data["error"] = f"Excel file not found: {abs_path}"
            return jsonify(response_data), 404

        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        workbook = excel.Workbooks.Open(abs_path, ReadOnly=False, UpdateLinks=False)

        ws_input  = get_ws_by_name(workbook, SHEETS[0])
        ws_sheet3 = get_ws_by_name(workbook, SHEETS[1])
        ws_sheet4 = get_ws_by_name(workbook, SHEETS[2])

        target_cells = TARGET_CELLS["Input data"][chiller_type]
        for i, cell in enumerate(target_cells):
            if i < len(values):
                try:
                    val = values[i]
                    ws_input.Range(cell).Value = float(val) if val not in (None, '') else 0.0
                except (ValueError, TypeError):
                    ws_input.Range(cell).Value = values[i]

        try:
            if REQUIRED_CAPACITY_INDEX < len(values):
                ws_sheet3.Range(TARGET_CELLS["Sheet3"][0]).Value = float(values[REQUIRED_CAPACITY_INDEX]) * KW_TO_TONR_FACTOR
        except (ValueError, TypeError): pass

        try:
            if EVAP_LEAVING_TEMP_INDEX < len(values):
                ws_sheet3.Range(SHEET3_ADJUSTED_CELLS["s_value_cell"]).Value = float(values[EVAP_LEAVING_TEMP_INDEX]) - 1.5
        except (ValueError, TypeError): pass

        try:
            input_idx = COND_LEAVING_TEMP_INDEX if chiller_type == "water" else AMBIENT_TEMP_INDEX
            if input_idx < len(values):
                ws_sheet3.Range(SHEET3_ADJUSTED_CELLS["d_value_cell"]).Value = float(values[input_idx]) + 2.0
        except (ValueError, TypeError): pass

        if chiller_type == "air":
            try:
                if AMBIENT_TEMP_INDEX  < len(values): ws_sheet3.Range(AIR_COOLED_CELLS["ambient_temp"]).Value  = float(values[AMBIENT_TEMP_INDEX])
                if ALTITUDE_INDEX      < len(values): ws_sheet3.Range(AIR_COOLED_CELLS["altitude"]).Value      = float(values[ALTITUDE_INDEX])
                if FIN_MATERIAL_INDEX  < len(values): ws_sheet3.Range(AIR_COOLED_CELLS["fin_material"]).Value  = str(values[FIN_MATERIAL_INDEX])
                if WET_BULB_TEMP_INDEX < len(values): ws_sheet3.Range(AIR_COOLED_CELLS["wet_bulb_temp"]).Value = float(values[WET_BULB_TEMP_INDEX])
            except (ValueError, TypeError): pass
        
        workbook.Save() 

        excel.Application.Run(f"'{os.path.basename(abs_path)}'!Iterate_Convergence")
        time.sleep(1) 
        excel.Calculate()
        time.sleep(1) 

        sheet4_res = {k: [safe_val(ws_sheet4.Range(c).Value) for c in v] for k, v in RESULT_COLUMN_CELLS["Sheet4"].items()}
        input_res  = {k: [safe_val(ws_input.Range(c).Value)  for c in v] for k, v in RESULT_COLUMN_CELLS["Input data"].items()}

        response_data["results"] = { "sheet4_results": sheet4_res, "input_data_results": input_res }

        label_map = {
            "Model": "D9", "Compressor": "D10",
            "Cooling/Heating Capacity (TR)": "E12", "Power Input (kW)": "E13",
            "Evaporator Inlet Temp (°C)": "E14",   "Evaporator Outlet Temp (°C)": "E15",
            "Evaporator Flow Rate (LPM)": "E16",    "Evaporator MEG (%)": "E17",
            "Evaporator Number of Passes (nos)": "E19", "Evaporator Fouling Factor (m²·K/kW)": "E20",
            "Evaporator Pressure Drop (kPa)": "N14", "Evaporator Tube Material": "E45",
            "Evaporator Saturation Temp (°C)": "N11", "Evaporator Velocity (m/s)": "N13",
            "IPLV/NPLV.IP_AHRIF": "E43"
        }

        for load, row in [(100,"31"),(90,"32"),(80,"33"),(75,"34"),(70,"35"),(60,"36"),(50,"37"),(40,"38"),(30,"39"),(25,"40"),(20,"41"),(10,"42")]:
            label_map[f"Part Load {load}% Capacity (TR)_AHRIF"]      = f"C{row}"
            label_map[f"Part Load {load}% Power (kW)_AHRIF"]         = f"D{row}"
            label_map[f"Part Load {load}% Efficiency (kW/TR)_AHRIF"] = f"E{row}"

        suffix = "_CCWET" if chiller_type == "water" else "_CAT"
        for load, row in [(100,"49"),(90,"50"),(80,"51"),(75,"52"),(70,"53"),(60,"54"),(50,"55"),(40,"56"),(30,"57"),(25,"58"),(20,"59"),(10,"60")]:
            label_map[f"Part Load {load}% Capacity (TR){suffix}"]      = f"C{row}"
            label_map[f"Part Load {load}% Power (kW){suffix}"]         = f"D{row}"
            label_map[f"Part Load {load}% Efficiency (kW/TR){suffix}"] = f"E{row}"

        if chiller_type == "water":
            label_map.update({
                "Condenser Water Inlet Temp (°C)": "E21", "Condenser Water Outlet Temp (°C)": "E22",
                "Condenser Flow Rate (LPM)": "E23",       "Condenser MEG (%)": "E24",
                "Condenser Number of Passes (nos)": "E26","Condenser Fouling Factor (m²·K/kW)": "E27",
                "Condenser Tube Material": "E44",         "Condenser Saturation Temp (°C)": "N12",
                "Condenser Velocity (m/s)": "N15",        "Condenser Side Pressure Drop (kPa)": "N16"
            })
        else: # air-cooled
            label_map.update({
                "Ambient Temperature (°C)": "E21", "Altitude (m)": "E28",
                "Fin Material": "E26",              "Wet Bulb Temperature (°C)": "E27",
                "No of Fan & Coil": "E22",          "Fan Type": "E23",
                "Cooling Type": "E24",              "Refrigerant charge": "E25",
                "Saturation Temp D (°C)": "E29"
            })

        det_results = {}
        for k, v in label_map.items():
            ws = ws_input if v.startswith('N') else ws_sheet4
            try: det_results[k] = safe_val(ws.Range(v).Value)
            except: det_results[k] = None

        try:
            det_results["Evaporator Model & Cu Tube Count"] = str(safe_val(ws_sheet4.Range("D18").Value)).strip()
            det_results["Condenser Model & Cu Tube Count"]  = str(safe_val(ws_sheet4.Range("D25").Value)).strip() if chiller_type == "water" else "N/A"
        except: pass

        response_data["detailed_results"] = det_results

    except Exception as e:
        logging.exception("Error in /write-excel")
        response_data["error"] = str(e)
        return jsonify(response_data), 500

    finally:
        try:
            if workbook: workbook.Close(SaveChanges=False)
            if excel:    excel.Quit()
        except Exception: pass
        pythoncom.CoUninitialize()

    return jsonify(response_data), 200

if __name__ == '__main__':
    logging.info("="*50)
    logging.info("  Kirloskar Chiller Selection Tool")
    logging.info("  Server running at: http://localhost:5000")
    logging.info("="*50)
    app.run(debug=True, port=5000)